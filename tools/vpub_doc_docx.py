# -*- coding: utf-8 -*-
"""Sinh docs/VPUB_VanPhongUBND/VPUB_MO_TA_CHI_TIEU.docx + .md (cùng nội dung, xem tools/doc_md.py) — tài liệu mô tả 5 chỉ tiêu Văn phòng UBND Thành phố trên IOC
(biểu đồ, trường, cách tính). Cùng khuôn với tools/vhtt_doc_docx.py.
Chạy từ thư mục gốc ioc:  python tools/vpub_doc_docx.py
Lời văn viết tay trong script; CÁC CON SỐ đọc thẳng từ js/core/vpub-config.js (qua node) nên luôn khớp giao diện.
Tên trường viết đúng như file mô tả đầu vào eform (tools/vpub-fields.json, trích từ sheet "Văn phòng UBND Thành phố Hà Nội").
Sau khi sinh, cập nhật Mục lục bằng Word: tools/vhtt_doc_finalize.ps1 -Docx <đường dẫn .docx> (dùng chung cho mọi tài liệu)."""
import json, os, re, subprocess
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from doc_md import MdMirror  # bản .md song song, cùng nội dung với .docx
from ctm_display import section as ctm_section  # phần "Cách hiển thị trên giao diện IOC", đọc từ cấu hình đang chạy

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIELDS = json.load(open(os.path.join(ROOT, 'tools', 'vpub-fields.json'), encoding='utf-8'))
OUT_DIR = os.path.join(ROOT, 'docs', 'VPUB_VanPhongUBND')
OUT = os.path.join(OUT_DIR, 'VPUB_MO_TA_CHI_TIEU.docx')
OUT_MD = os.path.join(OUT_DIR, 'VPUB_MO_TA_CHI_TIEU.md')
MD = MdMirror()

# ---------- Số liệu: chạy vpub-config.js bằng node, lấy VPUB_DATA / VPUB_SUM ----------
_JS = ("const fs=require('fs'),vm=require('vm');const c={};vm.createContext(c);"
       "vm.runInContext(fs.readFileSync('js/core/vpub-config.js','utf8')+';this.D=VPUB_DATA;this.S=VPUB_SUM;this.U=VPUB_UNITS;this.DV=VPUB_DV;',c);"
       "process.stdout.write(JSON.stringify({D:c.D,S:c.S,U:c.U,DV:c.DV}));")
_raw = subprocess.run(['node', '-e', _JS], cwd=ROOT, capture_output=True, check=True).stdout.decode('utf-8')
_J = json.loads(_raw)
D, S, U, DV = _J['D'], _J['S'], _J['U'], _J['DV']

def n(x, d=0):
    """Số kiểu Việt Nam: 1.234.567,8"""
    s = f'{abs(x):,.{d}f}'.replace(',', '#').replace('.', ',').replace('#', '.')
    return ('−' if x < 0 else '') + s
def pc(a, b, d=1): return n(a / b * 100, d) + '%'
def ssum(a): return sum(a)

Q3 = 2
s27, s29, s74, s52, s11 = S['cs27'], S['cs29'], S['cs74'], S['cs52'], S['cs1112']
d27, d29, d74, d11 = D['cs27'], D['cs29'], D['cs74'], D['cs1112']
h = s52['hs']
NDV = len(DV)

FONT = 'Times New Roman'
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2A, 0x5C, 0x9A)
GREY = RGBColor(0x55, 0x5F, 0x6D)
HEAD_FILL, KEY_FILL, NOTE_FILL, FORM_FILL = '1F3A5F', 'EEF3F9', 'FFF7E6', 'F3F6FA'
BORDER = 'A9B8CC'
PAGE_W = 16.0  # cm vùng chữ (A4, lề trái 3, phải 2)

doc = Document()

# ---------- Trang, kiểu chữ ----------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(2), Cm(2), Cm(2)
sec.header_distance, sec.footer_distance = Cm(1.1), Cm(1.1)

def set_font(style, size, bold=None, color=None, italic=None):
    style.font.name = FONT; style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'): rf.set(qn(a), FONT)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    if bold is not None: style.font.bold = bold
    if italic is not None: style.font.italic = italic
    if color is not None: style.font.color.rgb = color

st = doc.styles
set_font(st['Normal'], 13)
pf = st['Normal'].paragraph_format; pf.space_after = Pt(6); pf.line_spacing = 1.2
for name, size, before, after, color in (('Heading 1', 15, 18, 8, NAVY), ('Heading 2', 13.5, 14, 6, NAVY), ('Heading 3', 13, 10, 4, BLUE)):
    s = st[name]; set_font(s, size, bold=True, color=color, italic=False)
    s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True; s.paragraph_format.line_spacing = 1.15
for name in ('TOC 1', 'TOC 2', 'TOC 3'):
    try: s = st[name]
    except KeyError: s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    set_font(s, 12.5 if name == 'TOC 1' else 12, bold=(name == 'TOC 1'))
    s.paragraph_format.space_after = Pt(3); s.paragraph_format.left_indent = Cm({'TOC 1': 0, 'TOC 2': 0.6, 'TOC 3': 1.2}[name])
set_font(st['List Bullet'], 13); st['List Bullet'].paragraph_format.space_after = Pt(3)

# ---------- Tiện ích ----------
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill); tcPr.append(s)

def cell_margins(table, top=70, bottom=70, left=100, right=100):
    tblPr = table._tbl.tblPr; m = OxmlElement('w:tblCellMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement('w:' + k); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tblPr.append(m)

def borders(table, color=BORDER):
    tblPr = table._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for k in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + k); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), color); b.append(e)
    tblPr.append(b)

def fixed_layout(table, widths):
    tblPr = table._tbl.tblPr; lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(widths) * 567))); tw.set(qn('w:type'), 'dxa'); tblPr.append(tw)
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'), str(int(widths[i] * 567)))
    for row in table.rows:
        for i, c in enumerate(row.cells): c.width = Cm(widths[i])

def no_split(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement('w:cantSplit'); trPr.append(e)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement('w:tblHeader'); trPr.append(e)

def add_runs(p, text, size=None, bold=False, color=None, italic=False):
    """**đậm**, «tên trường» in nghiêng đậm màu xanh."""
    for part in re.split(r'(\*\*[^*]+\*\*|«[^»]+»)', text):
        if not part: continue
        if part.startswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('«'):
            r = p.add_run(part); r.bold = True; r.italic = True; r.font.color.rgb = BLUE
        else:
            r = p.add_run(part); r.bold = bold; r.italic = italic
            if color is not None: r.font.color.rgb = color
        if size: r.font.size = Pt(size)
    return p

def para(text, size=None, align=None, after=None, bold=False, color=None, italic=False, indent=None):
    p = _take_pb(doc.add_paragraph()); add_runs(p, text, size, bold, color, italic)
    MD.para(text, size, italic, bold)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if indent is not None: p.paragraph_format.left_indent = Cm(indent)
    return p

def keep_together(t):
    for row in t.rows[:-1]:
        for c in row.cells:
            for p in c.paragraphs: p.paragraph_format.keep_with_next = True

def table(rows, widths, header=True, size=11, key_col=False, fills=None, keep=False):
    MD.table(rows, header)
    t = doc.add_table(rows=len(rows), cols=len(widths)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t); cell_margins(t); fixed_layout(t, widths)
    for ri, row in enumerate(rows):
        no_split(t.rows[ri])
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
            lines = val if isinstance(val, list) else [val]
            for li, ln in enumerate(lines):
                if li: p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
                if header and ri == 0:
                    r = p.add_run(ln); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    add_runs(p, ln, size, bold=(key_col and ci == 0))
            if header and ri == 0: shade(c, HEAD_FILL)
            elif key_col and ci == 0: shade(c, KEY_FILL)
            elif fills and fills.get((ri, ci)): shade(c, fills[(ri, ci)])
    if header: repeat_header(t.rows[0])
    if keep: keep_together(t)
    gap = doc.add_paragraph(); gap.paragraph_format.space_after = Pt(4); gap.paragraph_format.line_spacing = 0.8
    return t

def box(lines, fill=FORM_FILL, label=None, size=12):
    MD.box(lines, label)
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, 'C9D4E3'); cell_margins(t, 110, 110, 160, 160); fixed_layout(t, [PAGE_W])
    no_split(t.rows[0])
    c = t.cell(0, 0); shade(c, fill); p = c.paragraphs[0]
    first = True
    if label:
        r = p.add_run(label); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = NAVY; first = False
    for ln in lines:
        if not first: p = c.add_paragraph()
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.15
        add_runs(p, ln, size); first = False
    gap = doc.add_paragraph(); gap.paragraph_format.space_after = Pt(4); gap.paragraph_format.line_spacing = 0.8

def field(p, instr):
    r = p.add_run(); b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._r.append(b)
    r = p.add_run(); t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = instr; r._r.append(t)
    r = p.add_run(); s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate'); r._r.append(s)
    r = p.add_run('…'); e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r2 = p.add_run(); r2._r.append(e)
    return r

# Sang trang = gắn "page_break_before" vào đoạn kế tiếp (tránh trang trắng khi đoạn ngắt rơi đúng cuối trang).
_PB = [False]
def page_break(): _PB[0] = True
def _take_pb(p):
    if _PB[0]: p.paragraph_format.page_break_before = True; _PB[0] = False
    return p
_orig_heading = doc.add_heading
def _heading(text, level=1):
    MD.heading(text, level)
    return _take_pb(_orig_heading(text, level=level))
doc.add_heading = _heading

def bottom_rule(p, color='1F3A5F', sz='8'):
    pPr = p._p.get_or_add_pPr(); bd = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), sz); b.set(qn('w:space'), '4'); b.set(qn('w:color'), color); bd.append(b); pPr.append(bd)

# ---------- Đầu trang, chân trang ----------
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hp.add_run('Văn phòng UBND Thành phố Hà Nội · Mô tả chỉ tiêu trên IOC'); r.italic = True; r.font.size = Pt(10); r.font.color.rgb = GREY
bottom_rule(hp, 'A9B8CC', '4')
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run('Trang '); r.font.size = Pt(10); r.font.color.rgb = GREY
r = field(fp, 'PAGE'); r.font.size = Pt(10)
r = fp.add_run(' / '); r.font.size = Pt(10); r.font.color.rgb = GREY
r = field(fp, 'NUMPAGES'); r.font.size = Pt(10)
sec.different_first_page_header_footer = True

N_FIELDS = sum(len(e['fields']) for e in FIELDS)

# ======================================================= TRANG BÌA
for _ in range(5): doc.add_paragraph()
para('HỆ THỐNG ĐIỀU HÀNH IOC/CAD', 13, WD_ALIGN_PARAGRAPH.CENTER, 0, bold=True)
para('THÀNH PHỐ HÀ NỘI', 13, WD_ALIGN_PARAGRAPH.CENTER, 0, bold=True)
p = para('', after=0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; bottom_rule(p, '1F3A5F', '6'); p.paragraph_format.left_indent = Cm(5.5); p.paragraph_format.right_indent = Cm(5.5)
for _ in range(4): doc.add_paragraph()
para('TÀI LIỆU MÔ TẢ CHỈ TIÊU', 22, WD_ALIGN_PARAGRAPH.CENTER, 6, bold=True, color=NAVY)
para('VĂN PHÒNG UBND THÀNH PHỐ HÀ NỘI', 20, WD_ALIGN_PARAGRAPH.CENTER, 14, bold=True, color=NAVY)
para('Biểu đồ, trường dữ liệu và cách tính số liệu của 5 chỉ tiêu trên IOC', 14, WD_ALIGN_PARAGRAPH.CENTER, 4, italic=True, color=GREY)
para('xây dựng từ 5 eform nhập liệu của Văn phòng', 14, WD_ALIGN_PARAGRAPH.CENTER, 0, italic=True, color=GREY)
for _ in range(6): doc.add_paragraph()
table([
    ['Đơn vị', 'Văn phòng Ủy ban nhân dân Thành phố Hà Nội'],
    ['Nguồn dữ liệu', f'File mô tả đầu vào eform, sheet "Văn phòng UBND Thành phố Hà Nội": 5 eform, {N_FIELDS} trường nhập'],
    ['Phạm vi', '5 chỉ tiêu (chỉ số 27, 29, 74, 52, 11-12), 17 biểu đồ, bảng 126 phường/xã'],
    ['Số liệu ví dụ', 'Số giả lập (chưa có số liệu thật), kỳ quý III và 9 tháng năm 2026'],
    ['Phiên bản', '1.6 · ngày 13/10/2026 (theo phản hồi nhóm triển khai khi nối eform thật: thêm quy ước xử lý dữ liệu eform — khóa bản ghi, cách lấy theo kỳ, chỉ hiện số có trường trong eform)'],
], [4.2, 10.0], header=False, size=12, key_col=True)
page_break()

# ======================================================= MỤC LỤC
para('MỤC LỤC', 15, WD_ALIGN_PARAGRAPH.CENTER, 10, bold=True, color=NAVY)
field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u'); MD.toc()
page_break()

# ======================================================= PHẦN I
doc.add_heading('PHẦN I. QUY ƯỚC CHUNG', level=1)
doc.add_heading('1. Mục đích tài liệu', level=2)
para('Tài liệu mô tả cách IOC dựng 5 chỉ tiêu của Văn phòng UBND Thành phố từ 5 eform nhập liệu: mỗi chỉ tiêu có những biểu đồ nào, dùng những trường nào của eform, mỗi trường dùng để làm gì, và mỗi con số được cộng, đếm, tính toán ra sao. Tên trường viết đúng như trong file mô tả đầu vào eform để cán bộ nhập liệu và đội kỹ thuật cùng đối chiếu.')
SHOW_NOTE_BOX = False  # khối "Lưu ý, đề nghị Văn phòng xác nhận" cuối mỗi chỉ tiêu — ẩn theo yêu cầu 14/10/2026; nội dung vẫn gom vào Phần IV
para('Cuối mỗi chỉ tiêu có mục **"Lưu ý, đề nghị Văn phòng xác nhận"**; Phần IV gom toàn bộ các đề nghị đó vào một bảng để tiện rà soát.' if SHOW_NOTE_BOX
     else 'Các đề nghị Văn phòng xác nhận được gom vào một bảng ở Phần IV để tiện rà soát.')
box(['Mọi con số ví dụ trong tài liệu là **số giả lập** dùng để dựng giao diện, đã được cộng kiểm cho khớp (tổng luôn bằng cộng các phần, số toàn Thành phố bằng cộng 126 phường/xã). Khi có số liệu thật, cách tính giữ nguyên, chỉ thay số.'], NOTE_FILL, 'Lưu ý: ')

doc.add_heading('2. Kỳ báo cáo', level=2)
para('Mọi eform đều có 4 trường xác định kỳ: «Kỳ báo cáo» (Tháng / Quý / Năm), «Năm», «Quý», «Tháng». IOC dùng 4 trường này để chọn bản ghi thuộc kỳ đang xem. Có hai cách lấy dữ liệu theo kỳ:')
table([
    ['Cách lấy', 'Áp dụng cho', 'Ý nghĩa'],
    ['**Bản của kỳ đang xem**', 'Chỉ tiêu 1, 2, 5', 'Số trong eform là số tại một thời điểm (tỷ lệ của quý, tình trạng triển khai trợ lý ảo, số thiết bị đang có). Chỉ lấy bản báo cáo của kỳ đang xem, không cộng các kỳ với nhau.'],
    ['**Cộng dồn**', 'Chỉ tiêu 3, 4', 'Chỉ tiêu 3: lấy mọi nhiệm vụ có «Ngày giao» trong kỳ (mặc định 01/01 – 30/9/2026), trạng thái theo bản cập nhật mới nhất. Chỉ tiêu 4: cộng số liệu của quý I, II và III.'],
], [3.6, 3.0, 9.4])

doc.add_heading('3. Các phép tính dùng trong tài liệu', level=2)
table([
    ['Ký hiệu', 'Ý nghĩa', 'Ví dụ'],
    ['ĐẾM(điều kiện)', 'Đếm số bản ghi thỏa điều kiện', 'ĐẾM(«Trạng thái thực hiện» = Hoàn thành đúng hạn)'],
    ['TỔNG(«trường»)', 'Cộng giá trị của trường trên các bản ghi được lấy', 'TỔNG(«Tổng số hồ sơ tiếp nhận»)'],
    ['Nhóm theo «trường»', 'Chia bản ghi thành từng nhóm theo giá trị của trường, rồi đếm / cộng trong từng nhóm', 'Nhóm theo «Cơ quan giao»'],
    ['Tỷ lệ (%)', 'Tử số ÷ Mẫu số × 100, làm tròn 1 chữ số thập phân', f'{n(s74["ht"])} ÷ {n(s74["tong"])} × 100 = {pc(s74["ht"], s74["tong"])}'],
    ['Bình quân có trọng số', 'TỔNG(giá trị × trọng số) ÷ TỔNG(trọng số); đơn vị lớn ảnh hưởng nhiều hơn đơn vị nhỏ', 'Mức độ hài lòng («Điểm hài lòng (%)»), trọng số là số hồ sơ tiếp nhận'],
    ['Lũy kế', 'Cộng dồn từ tháng 1 đến tháng t', 'Lũy kế T3 = T1 + T2 + T3'],
    ['Khóa chống trùng', 'Trường dùng để nhận ra một bản ghi; bản ghi trùng khóa chỉ tính một lần (lấy bản mới nhất)', '«Mã nhiệm vụ»'],
], [3.6, 6.6, 5.8])

doc.add_heading('4. Vai trò của trường trong cách tính', level=2)
para('Ở mỗi chỉ tiêu, bảng "Trường dữ liệu sử dụng" ghi rõ vai trò của từng trường:')
table([
    ['Vai trò', 'Giải thích'],
    ['Lọc kỳ', 'Chọn bản ghi thuộc kỳ đang xem.'],
    ['Khóa', 'Nhận diện một bản ghi, dùng để đếm và chống đếm trùng.'],
    ['Điều kiện', 'Quyết định bản ghi có được tính hay không, hoặc thuộc nhóm nào.'],
    ['Nhóm', 'Tạo các cột / lát cắt của biểu đồ (trục danh mục).'],
    ['Cộng', 'Giá trị số được cộng lại để ra con số hiển thị.'],
], [3.6, 12.4], key_col=True)
para('Các trường không có vai trò trong biểu đồ (tên, văn bản, ghi chú…) vẫn được lưu và hiện ở danh sách bản ghi chi tiết; tài liệu liệt kê ở cuối bảng trường của từng chỉ tiêu.', italic=True, color=GREY, size=12)

doc.add_heading('5. Nguyên tắc tính tỷ lệ', level=2)
para('Nhiều eform có sẵn trường tỷ lệ do đơn vị tự tính (ví dụ «Tỷ lệ (%) (tự tính)» của chỉ số 27). IOC **không lấy trung bình** các tỷ lệ này, vì mỗi đơn vị có quy mô khác nhau. IOC cộng tử số và mẫu số của tất cả đơn vị rồi mới chia.')
u0 = U[9]  # Phường Hoàn Kiếm
box([f'{u0["ten"]} quý III: {n(u0["vbMang"][Q3])} ÷ {n(u0["vbTong"][Q3])} = {pc(u0["vbMang"][Q3], u0["vbTong"][Q3])}. '
     f'Tỷ lệ chung của 126 phường/xã không phải trung bình của 126 tỷ lệ, mà là {n(s27["xaMang"][Q3])} ÷ {n(s27["xaTong"][Q3])} = {pc(s27["xaMang"][Q3], s27["xaTong"][Q3])}.'],
    FORM_FILL, 'Ví dụ: ')

doc.add_heading('6. Đơn vị báo cáo', level=2)
para(f'Danh mục «Đơn vị báo cáo» (và «Địa bàn») trong cả 5 eform hiện gồm đúng **{NDV} phường/xã** ({sum(1 for x in DV if x.startswith("Phường"))} phường, {sum(1 for x in DV if x.startswith("Xã"))} xã). Hai eform chỉ số 27 và 11-12 có thêm trường «Cấp báo cáo» để tách số của các sở, ban, ngành. Trong số giả lập, phần sở, ban, ngành là số cộng chung của {d27["soDv"]} đơn vị, chưa tách từng sở vì danh mục chưa có tên các sở.')

doc.add_heading('7. Tổng hợp 5 chỉ tiêu', level=2)
table([
    ['#', 'Chỉ tiêu trên IOC', 'Eform nguồn', 'Kỳ hiển thị', 'Biểu đồ'],
    ['1', 'Tỷ lệ văn bản, hồ sơ công việc xử lý trên môi trường mạng (chỉ số 27)', FIELDS[0]['eform'], 'Quý III/2026', '3'],
    ['2', 'Triển khai trợ lý ảo (chỉ số 29)', FIELDS[1]['eform'], 'Năm 2026', '3'],
    ['3', 'Theo dõi nhiệm vụ Chính phủ, Thủ tướng Chính phủ giao (chỉ số 74)', FIELDS[2]['eform'], '9 tháng 2026', '4'],
    ['4', 'Hồ sơ thủ tục hành chính, dịch vụ công (chỉ số 52)', FIELDS[3]['eform'], '9 tháng 2026', '3'],
    ['5', 'Hạ tầng công nghệ thông tin (chỉ số 11, 12)', FIELDS[4]['eform'], 'Quý III/2026', '4'],
], [0.8, 5.6, 5.6, 2.4, 1.6], size=10.5)
para('Mỗi chỉ tiêu trên IOC gồm: 1 khối **thẻ số tổng hợp** và các **biểu đồ**; cách tính từng số chỉ trình bày trong tài liệu này, không hiện trên giao diện. Bên phải là bảng **126 phường/xã** (Phần III). Bấm vào một biểu đồ thì bảng số của chính biểu đồ đó hiện ở ô bên phải.')
para('Bộ chỉ tiêu cũ của Văn phòng (1 chỉ tiêu STT 74 lấy qua API, 4.860 nhiệm vụ lũy kế) đã được thay bằng bộ này và đã xóa khỏi IOC.', italic=True, color=GREY, size=12)

# ======================================================= PHẦN II
page_break()
doc.add_heading('PHẦN II. CHI TIẾT 5 CHỈ TIÊU', level=1)
para('Mỗi chỉ tiêu được trình bày theo cùng một khuôn: (1) Thông tin chung · (2) Trường dữ liệu sử dụng · (3) Thẻ số tổng hợp · (4) Từng biểu đồ' + (' · (5) Lưu ý, đề nghị Văn phòng xác nhận' if SHOW_NOTE_BOX else '') + '.', italic=True, color=GREY)

USE_W = [0.8, 5.4, 2.0, 2.3, 5.5]
def fields_table(rows, unused):
    data = [['#', 'Tên trường (đúng như eform)', 'Kiểu nhập', 'Vai trò', 'Dùng để']]
    for i, r in enumerate(rows): data.append([str(i + 1), r[0], r[1], r[2], r[3]])
    table(data, USE_W, size=10.5)
    if unused:
        para('**Trường không dùng trong biểu đồ** (chỉ hiện ở danh sách chi tiết): ' + ', '.join('«' + u + '»' for u in unused) + '.', size=11.5, color=GREY)

def kpi_table(rows):
    table([['Thẻ số', 'Công thức', 'Kết quả (giả lập)']] + rows, [3.6, 8.2, 4.2], size=11, keep=True)

def chart(code, title, kind, axis, value, formula, example, note=None):
    doc.add_heading(f'Biểu đồ {code}. {title}', level=3)
    rows = [['Loại biểu đồ', kind], ['Trục / nhóm', axis], ['Giá trị hiển thị', value], ['Cách tính', formula], ['Ví dụ (giả lập)', example]]
    if note: rows.append(['Ghi chú', note])
    table(rows, [3.4, 12.6], header=False, size=11.5, key_col=True, keep=True)

def info_table(rows):
    table(rows, [4.0, 12.0], header=False, size=11.5, key_col=True, keep=True)

ALL_NOTES = []
def notes(ct, items):
    if SHOW_NOTE_BOX: box(['• ' + x for x in items], NOTE_FILL, 'Lưu ý, đề nghị Văn phòng xác nhận', 12)
    ALL_NOTES.extend((ct, x) for x in items)

def ky(txt): return [[['Kỳ báo cáo, Năm,', 'Quý, Tháng'], 'Danh mục / Số', 'Lọc kỳ', txt]]

# ---------------- CHỈ TIÊU 1 — chỉ số 27
doc.add_heading('Chỉ tiêu 1. Tỷ lệ văn bản, hồ sơ công việc xử lý trên môi trường mạng (chỉ số 27)', level=2)
doc.add_heading('1.1. Thông tin chung', level=3)
info_table([
    ['Eform nguồn', FIELDS[0]['eform'].strip()],
    ['Một bản ghi là', 'Số liệu một quý của một đơn vị báo cáo (một sở, ban, ngành hoặc một phường/xã)'],
    ['Cách lấy theo kỳ', '**Bản của kỳ đang xem.** Lấy các bản ghi của quý đang xem, không cộng các quý với nhau.'],
    ['Kỳ hiển thị', f'Quý III/2026 · {d27["soDv"] + NDV} đơn vị ({d27["soDv"]} sở, ban, ngành và {NDV} phường/xã)'],
])
doc.add_heading('1.2. Trường dữ liệu sử dụng', level=3)
fields_table(ky('Chọn các bản ghi của quý đang xem; «Quý» còn là trục của biểu đồ 1.1') + [
    ['Đơn vị báo cáo', 'Danh mục', 'Nhóm', 'Tính tỷ lệ của từng phường/xã: biểu đồ 1.2, 1.3 và bảng phường/xã'],
    ['Cấp báo cáo', 'Danh mục', 'Nhóm', 'Tách Sở, ban, ngành / Phường/xã: thẻ số theo cấp, 2 đoạn cột của biểu đồ 1.1'],
    ['Tổng số văn bản, hồ sơ công việc', 'Số', 'Cộng', 'Mẫu số của mọi tỷ lệ'],
    ['Số văn bản, hồ sơ xử lý trên môi trường mạng', 'Số', 'Cộng', 'Tử số của mọi tỷ lệ'],
], ['Tỷ lệ (%) (tự tính)', 'Ghi chú'])
para('«Tỷ lệ (%) (tự tính)» không được dùng vì IOC tự tính lại tỷ lệ từ hai trường số (xem Phần I, mục 5).', size=11.5, color=GREY, italic=True)
doc.add_heading('1.3. Thẻ số tổng hợp', level=3)
kpi_table([
    ['Tổng văn bản, hồ sơ', 'TỔNG(«Tổng số văn bản, hồ sơ công việc») của quý III, cả 2 cấp', n(s27['tong'][Q3])],
    ['Xử lý trên môi trường mạng', 'TỔNG(«Số văn bản, hồ sơ xử lý trên môi trường mạng») của quý III, cả 2 cấp', n(s27['mang'][Q3])],
    ['Tỷ lệ toàn Thành phố', 'Xử lý trên môi trường mạng ÷ Tổng văn bản, hồ sơ × 100; dòng phụ là tỷ lệ quý trước', f'{pc(s27["mang"][Q3], s27["tong"][Q3])} (quý II: {pc(s27["mang"][1], s27["tong"][1])})'],
    ['Sở, ban, ngành', 'Như trên, chỉ lấy «Cấp báo cáo» = Cấp Sở, ban, ngành', f'{n(d27["soMang"][Q3])} ÷ {n(d27["soTong"][Q3])} = {pc(d27["soMang"][Q3], d27["soTong"][Q3])}'],
    ['Phường/xã', 'Như trên, chỉ lấy «Cấp báo cáo» = Xã/phường', f'{n(s27["xaMang"][Q3])} ÷ {n(s27["xaTong"][Q3])} = {pc(s27["xaMang"][Q3], s27["xaTong"][Q3])}'],
])
doc.add_heading('1.4. Biểu đồ', level=3)
QN = ['Quý I', 'Quý II', 'Quý III']
chart('1.1', 'Văn bản, hồ sơ và tỷ lệ xử lý trên mạng theo quý', 'Biểu đồ kết hợp: cột xếp chồng (trục trái) + đường (trục phải)',
      'Trục ngang: «Quý» I → III; mỗi cột chia 2 đoạn theo «Cấp báo cáo»',
      'Cột: số văn bản, hồ sơ · Đường: tỷ lệ xử lý trên mạng toàn Thành phố (%)',
      ['Nhóm theo «Quý» và «Cấp báo cáo»:', 'Mỗi đoạn cột = TỔNG(«Tổng số văn bản, hồ sơ công việc»);',
       'Đường = TỔNG(«Số văn bản, hồ sơ xử lý trên môi trường mạng») ÷ TỔNG(«Tổng số văn bản, hồ sơ công việc») × 100, tính chung 2 cấp.'],
      ' · '.join(f'{QN[q]}: {n(s27["tong"][q])} văn bản, hồ sơ, {pc(s27["mang"][q], s27["tong"][q])}' for q in range(3)))
nh = d27['nhom']
chart('1.2', 'Phường/xã theo mức tỷ lệ — quý III', 'Biểu đồ tròn dạng vành khuyên; giữa vòng ghi số phường/xã',
      '4 nhóm: ' + ' / '.join(g[0] for g in nh), 'Số phường/xã của từng nhóm; di chuột hiện thêm tỷ trọng %',
      'Với mỗi phường/xã: tỷ lệ = «Số văn bản, hồ sơ xử lý trên môi trường mạng» ÷ «Tổng số văn bản, hồ sơ công việc» × 100; xếp vào nhóm theo ngưỡng 95% / 90% / 80%, rồi đếm số đơn vị.',
      ' · '.join(f'{g[0]} {s27["nhom"][i]}' for i, g in enumerate(nh)) + f' → giữa vòng {NDV}',
      'Ngưỡng 95% / 90% / 80% do IOC tự đặt để xem phân bố, chưa phải ngưỡng chính thức.')
th = s27['thapNhat']
chart('1.3', '10 phường/xã có tỷ lệ thấp nhất — quý III', 'Biểu đồ cột ngang',
      'Trục dọc: 10 phường/xã, xếp từ tỷ lệ thấp nhất', 'Tỷ lệ xử lý trên môi trường mạng (%)',
      'Tính tỷ lệ của từng phường/xã như biểu đồ 1.2, xếp tăng dần, lấy 10 đơn vị đầu.',
      ' · '.join(f'{x[0]} {n(x[1], 1)}%' for x in th[:3]) + f' … {th[-1][0]} {n(th[-1][1], 1)}%',
      'Dùng để biết đơn vị nào cần đôn đốc trước.')
notes('1', ['Danh mục «Cấp báo cáo» hiện chỉ có một giá trị "Cấp Sở, ban, ngành". Đề nghị thêm "Xã/phường" như eform chỉ số 11, 12.',
            f'Danh mục «Đơn vị báo cáo» hiện chỉ có {NDV} phường/xã. Đề nghị bổ sung các sở, ban, ngành để tách được số của từng sở.',
            '«Tỷ lệ (%) (tự tính)» trùng với số IOC tự tính; đề nghị bỏ trường này hoặc để chỉ đọc, tránh đơn vị nhập tỷ lệ không khớp với hai trường số (dữ liệu mẫu đang ghi 10.000.000 / 10.000.000 nhưng tỷ lệ 32).',
            'Ngưỡng nhóm 95% / 90% / 80% ở biểu đồ 1.2 cần Văn phòng xác nhận hoặc thay bằng chỉ tiêu giao.'])

# ---------------- CHỈ TIÊU 2 — chỉ số 29
page_break()
doc.add_heading('Chỉ tiêu 2. Triển khai trợ lý ảo (chỉ số 29)', level=2)
doc.add_heading('2.1. Thông tin chung', level=3)
info_table([
    ['Eform nguồn', FIELDS[1]['eform']],
    ['Một bản ghi là', 'Tình trạng triển khai trợ lý ảo của một đơn vị báo cáo trong năm'],
    ['Cách lấy theo kỳ', '**Bản của kỳ đang xem.** Lấy bản báo cáo mới nhất của mỗi đơn vị trong năm đang xem.'],
    ['Kỳ hiển thị', f'Năm 2026 · {s29["dv"]} phường/xã'],
])
doc.add_heading('2.2. Trường dữ liệu sử dụng', level=3)
fields_table(ky('Chọn bản báo cáo của năm đang xem') + [
    ['Đơn vị báo cáo', 'Danh mục', 'Khóa, Nhóm', 'Mỗi đơn vị một bản ghi; tách Phường / Xã theo tên đơn vị cho biểu đồ 2.1'],
    ['Đã triển khai trợ lý ảo', '(chưa ghi)', 'Điều kiện', 'Có / Chưa: thẻ số, biểu đồ 2.1; chỉ đơn vị "Có" mới vào biểu đồ 2.2, 2.3'],
    ['Tên ứng dụng trợ lý ảo', 'Chữ', 'Nhóm', 'Gom về 3 nhóm ứng dụng: biểu đồ 2.2'],
    ['Đối tượng phục vụ', 'Chữ', 'Nhóm', 'Lát cắt của biểu đồ 2.3'],
], ['Ghi chú'])
doc.add_heading('2.3. Thẻ số tổng hợp', level=3)
kpi_table([
    ['Đơn vị báo cáo', 'ĐẾM(«Đơn vị báo cáo») đã gửi báo cáo của năm', f'{n(s29["dv"])} đơn vị'],
    ['Đã triển khai trợ lý ảo', 'ĐẾM(«Đã triển khai trợ lý ảo» = Có)', f'{n(s29["tla"])} đơn vị'],
    ['Tỷ lệ đã triển khai', 'Đã triển khai ÷ Đơn vị báo cáo × 100', f'{n(s29["tla"])} ÷ {n(s29["dv"])} = {pc(s29["tla"], s29["dv"])}'],
    ['Chưa triển khai', 'Đơn vị báo cáo − Đã triển khai', f'{n(s29["dv"])} − {n(s29["tla"])} = {n(s29["dv"] - s29["tla"])}'],
])
doc.add_heading('2.4. Biểu đồ', level=3)
ph, xa = s29['phuong'], s29['xa']
chart('2.1', 'Phường và xã', 'Biểu đồ cột ngang xếp chồng',
      'Trục dọc: Phường / Xã (theo tên trong «Đơn vị báo cáo»); mỗi cột chia 2 đoạn Đã triển khai / Chưa triển khai',
      'Số đơn vị; số ở đầu cột = tổng số phường (hoặc xã)',
      'Tách đơn vị theo chữ đầu tên (Phường / Xã), trong mỗi nhóm ĐẾM theo «Đã triển khai trợ lý ảo».',
      f'Phường: {ph[0]} + {ph[1] - ph[0]} = {ph[1]} ({pc(ph[0], ph[1])}) · Xã: {xa[0]} + {xa[1] - xa[0]} = {xa[1]} ({pc(xa[0], xa[1])})')
chart('2.2', 'Theo ứng dụng trợ lý ảo', 'Biểu đồ tròn dạng vành khuyên; giữa vòng ghi số đơn vị đã triển khai',
      '3 nhóm ứng dụng: ' + ' / '.join(a[0] for a in d29['app']), 'Số đơn vị',
      'Chỉ lấy đơn vị «Đã triển khai trợ lý ảo» = Có; IOC gom «Tên ứng dụng trợ lý ảo» về 3 nhóm rồi ĐẾM.',
      ' · '.join(f'{a[0]} {s29["app"][i]}' for i, a in enumerate(d29['app'])) + f' → giữa vòng {s29["tla"]}',
      '«Tên ứng dụng trợ lý ảo» là chữ gõ tự do, nên việc gom nhóm cần một danh sách quy đổi tên do Văn phòng duyệt.')
chart('2.3', 'Theo đối tượng phục vụ', 'Biểu đồ tròn dạng vành khuyên; giữa vòng ghi số đơn vị đã triển khai',
      'Từng giá trị của «Đối tượng phục vụ»: ' + ' / '.join(a[0] for a in d29['doiTuong']), 'Số đơn vị',
      'Chỉ lấy đơn vị đã triển khai; nhóm theo «Đối tượng phục vụ», mỗi lát = ĐẾM(«Đơn vị báo cáo»).',
      ' · '.join(f'{a[0]} {s29["dt"][i]}' for i, a in enumerate(d29['doiTuong'])) + f' → giữa vòng {s29["tla"]}')
notes('2', ['«Đã triển khai trợ lý ảo» chưa có kiểu nhập trong file mô tả. Đề nghị đặt là Danh mục với 2 giá trị Có / Chưa.',
            '«Tên ứng dụng trợ lý ảo» đang là Chữ. Đề nghị đổi sang Danh mục các ứng dụng đang dùng (thêm giá trị "Khác") để thống kê chính xác.',
            '«Đối tượng phục vụ» đang là Chữ; eform chỉ số 11, 12 đã có danh mục cùng tên (Công chức, viên chức / Người dân và doanh nghiệp). Đề nghị dùng lại danh mục đó, thêm giá trị "Cả hai".'])

# ---------------- CHỈ TIÊU 3 — chỉ số 74
page_break()
doc.add_heading('Chỉ tiêu 3. Theo dõi nhiệm vụ Chính phủ, Thủ tướng Chính phủ giao (chỉ số 74)', level=2)
doc.add_heading('3.1. Thông tin chung', level=3)
info_table([
    ['Eform nguồn', FIELDS[2]['eform']],
    ['Một bản ghi là', 'Một nhiệm vụ («Mã nhiệm vụ»)'],
    ['Cách lấy theo kỳ', '**Cộng dồn** theo «Ngày giao»: lấy mọi nhiệm vụ giao từ 01/01 đến hết kỳ đang xem; trạng thái lấy theo bản cập nhật mới nhất; mỗi «Mã nhiệm vụ» chỉ tính một lần.'],
    ['Kỳ hiển thị', f'9 tháng năm 2026 · {n(s74["tong"])} nhiệm vụ · trạng thái tại 30/9/2026'],
])
doc.add_heading('3.2. Trường dữ liệu sử dụng', level=3)
fields_table(ky('Chọn bản cập nhật mới nhất') + [
    ['Mã nhiệm vụ', 'Chữ', 'Khóa', 'Đếm số nhiệm vụ, chống đếm trùng'],
    ['Cơ quan giao', 'Chữ', 'Nhóm', 'Trục của biểu đồ 3.2'],
    ['Ngày giao', 'Ngày tháng', 'Lọc kỳ, Nhóm', 'Chọn nhiệm vụ thuộc kỳ; đường "Đã giao" của biểu đồ 3.4'],
    ['Hạn hoàn thành', 'Ngày tháng', 'Điều kiện', 'Xác định nhiệm vụ sắp đến hạn, quá hạn'],
    ['Đơn vị chủ trì', 'Chữ', 'Nhóm', 'Trục của biểu đồ 3.3'],
    ['Trạng thái thực hiện', 'Danh mục', 'Điều kiện, Nhóm', 'Thẻ số; biểu đồ 3.1, 3.2, 3.3'],
    ['Tỷ lệ hoàn thành (%)', 'Số', 'Cộng', 'Thẻ Tiến độ bình quân'],
    ['Ngày hoàn thành', 'Ngày tháng', 'Nhóm', 'Đường "Đã hoàn thành" của biểu đồ 3.4'],
    ['Cảnh báo tiến độ', 'Chữ', 'Điều kiện', 'Thẻ Sắp đến hạn'],
    ['Địa bàn', 'Danh mục', 'Nhóm', 'Cột "Nhiệm vụ CP, TTg có địa bàn" của bảng phường/xã'],
], ['Tên nhiệm vụ', 'Văn bản giao nhiệm vụ', 'Đơn vị phối hợp', 'Ghi chú'])
para('«Cơ quan giao» là cấp giao nhiệm vụ cho Thành phố (Chính phủ, Thủ tướng…); «Đơn vị chủ trì» là đơn vị của Thành phố chịu trách nhiệm thực hiện. Hai trường này khác nhau và được thống kê riêng. «Đơn vị phối hợp» không được đếm vào biểu đồ.', size=11.5, color=GREY, italic=True)
doc.add_heading('3.3. Trạng thái thực hiện (dùng chung cho các biểu đồ)', level=3)
para('Danh mục «Trạng thái thực hiện» của eform có 6 giá trị. IOC dùng nguyên 6 giá trị, không tự phân loại lại.')
ST = [x['l'] for x in d74['st']]
MEAN = ['Đã xong, «Ngày hoàn thành» ≤ «Hạn hoàn thành»', 'Đã xong, «Ngày hoàn thành» > «Hạn hoàn thành»', 'Đang làm, chưa đến hạn',
        'Chưa bắt đầu, chưa đến hạn', 'Đã qua «Hạn hoàn thành» mà chưa xong', 'Dừng theo chỉ đạo, không tính quá hạn']
table([['Trạng thái', 'Ý nghĩa', 'Số NV']] + [[ST[i], MEAN[i], n(s74['st'][i])] for i in range(6)] + [['**Cộng**', '', f'**{n(s74["tong"])}**']],
      [4.6, 9.7, 1.7], size=11)
doc.add_heading('3.4. Thẻ số tổng hợp', level=3)
kpi_table([
    ['Tổng số nhiệm vụ', 'ĐẾM(«Mã nhiệm vụ») có «Ngày giao» trong kỳ', f'{n(s74["tong"])} nhiệm vụ'],
    ['Đã hoàn thành', 'Hoàn thành đúng hạn + Hoàn thành quá hạn', f'{n(s74["st"][0])} + {n(s74["st"][1])} = {n(s74["ht"])}'],
    ['Tỷ lệ đúng hạn', 'Hoàn thành đúng hạn ÷ Đã hoàn thành × 100 (dòng phụ dưới thẻ Đã hoàn thành)', f'{n(s74["st"][0])} ÷ {n(s74["ht"])} = {pc(s74["st"][0], s74["ht"])}'],
    ['Tỷ lệ hoàn thành', 'Đã hoàn thành ÷ Tổng số nhiệm vụ × 100', f'{n(s74["ht"])} ÷ {n(s74["tong"])} = {pc(s74["ht"], s74["tong"])}'],
    ['Quá hạn chưa hoàn thành', 'ĐẾM(«Trạng thái thực hiện» = Quá hạn chưa hoàn thành)', f'{n(s74["st"][4])} nhiệm vụ'],
    ['Sắp đến hạn', 'ĐẾM(nhiệm vụ chưa hoàn thành có «Cảnh báo tiến độ» = sắp đến hạn, tức còn ≤ 7 ngày đến «Hạn hoàn thành»)', f'{n(d74["sapDenHan"])} nhiệm vụ'],
    ['Tiến độ bình quân', 'Trung bình «Tỷ lệ hoàn thành (%)» của các nhiệm vụ Đang thực hiện', n(d74['tienDoBq'], 1) + '%'],
])
doc.add_heading('3.5. Biểu đồ', level=3)
chart('3.1', 'Nhiệm vụ theo trạng thái thực hiện', 'Biểu đồ tròn dạng vành khuyên; giữa vòng ghi tổng số nhiệm vụ',
      '6 giá trị của «Trạng thái thực hiện» (mục 3.3)', 'Số nhiệm vụ; di chuột hiện thêm tỷ trọng %',
      'Nhóm theo «Trạng thái thực hiện», mỗi lát = ĐẾM(«Mã nhiệm vụ»).',
      ' · '.join(f'{ST[i]} {n(s74["st"][i])}' for i in range(6)) + f' → giữa vòng {n(s74["tong"])}')
cq = d74['coQuan']
chart('3.2', 'Nhiệm vụ theo cơ quan giao', 'Biểu đồ cột ngang xếp chồng',
      'Trục dọc: từng giá trị của «Cơ quan giao»; mỗi cột chia 6 đoạn màu theo trạng thái',
      'Số trong đoạn = số nhiệm vụ; số ở đầu cột = tổng nhiệm vụ của cơ quan (đoạn quá hẹp không ghi số)',
      'Nhóm theo «Cơ quan giao», trong mỗi nhóm ĐẾM(«Mã nhiệm vụ») theo «Trạng thái thực hiện».',
      ' · '.join(f'{r[0]} {n(ssum(r[1]))}' for r in cq) + f' (cộng {n(s74["tong"])}). Ví dụ {cq[1][0]}: ' + ' + '.join(n(v) for v in cq[1][1]) + f' = {n(ssum(cq[1][1]))}.')
ct = d74['chuTri']; g4 = [g[0] for g in d74['nhom4']]
chart('3.3', 'Nhiệm vụ theo đơn vị chủ trì', 'Biểu đồ cột ngang xếp chồng',
      'Trục dọc: «Đơn vị chủ trì» (các sở có nhiều nhiệm vụ nhất, còn lại gộp "Các đơn vị khác"); mỗi cột chia 4 đoạn: ' + ' / '.join(g4),
      'Số nhiệm vụ; số ở đầu cột = tổng nhiệm vụ đơn vị chủ trì',
      ['Nhóm theo «Đơn vị chủ trì», gộp 6 trạng thái thành 4 nhóm:', 'Đã hoàn thành = Hoàn thành đúng hạn + Hoàn thành quá hạn;',
       'Đang/chưa thực hiện, trong hạn = Đang thực hiện + Chưa thực hiện;', 'Quá hạn chưa hoàn thành; Tạm dừng.'],
      f'{ct[0][0]}: ' + ' + '.join(n(v) for v in ct[0][1]) + f' = {n(ssum(ct[0][1]))}. Cộng {len(ct)} nhóm đơn vị = {n(s74["tong"])}.',
      'Mỗi nhiệm vụ chỉ tính cho một đơn vị chủ trì; đơn vị phối hợp không được đếm.')
lg, lh = s74['lkGiao'], s74['lkHt']
chart('3.4', 'Giao và hoàn thành lũy kế theo tháng', 'Biểu đồ đường, 2 đường: Đã giao và Đã hoàn thành',
      'Trục ngang: tháng T1 → T9; trục dọc: số nhiệm vụ', 'Số nhiệm vụ lũy kế đến cuối mỗi tháng',
      ['Đã giao(t) = ĐẾM(«Ngày giao» ≤ ngày cuối tháng t).', 'Đã hoàn thành(t) = ĐẾM(«Ngày hoàn thành» ≤ ngày cuối tháng t).'],
      f'T1: {n(lg[0])} / {n(lh[0])} · T6: {n(lg[5])} / {n(lh[5])} · T9: {n(lg[8])} / {n(lh[8])}. Khoảng cách tháng 9 = {n(lg[8] - lh[8])} nhiệm vụ chưa hoàn thành.',
      'Khoảng cách giữa hai đường là số nhiệm vụ đang tồn; đường càng xa nhau thì tồn đọng càng lớn.')
notes('3', ['«Cơ quan giao» và «Đơn vị chủ trì» đang là Chữ gõ tự do. Đề nghị đổi sang Danh mục để IOC gom nhóm chính xác (Chính phủ / Thủ tướng Chính phủ / Phó Thủ tướng Chính phủ; danh sách sở, ngành, phường/xã).',
            '«Cảnh báo tiến độ» đang là Chữ và chưa có giá trị mẫu. Đề nghị đổi sang Danh mục (Bình thường / Sắp đến hạn / Quá hạn) hoặc để IOC tự tính từ «Hạn hoàn thành». IOC đang coi còn ≤ 7 ngày là sắp đến hạn.',
            'Eform không có trường «Đơn vị báo cáo»; IOC hiểu là Văn phòng nhập tập trung toàn bộ nhiệm vụ. Đề nghị xác nhận.',
            f'Bộ chỉ tiêu cũ (4.860 nhiệm vụ lũy kế qua API) đã thay bằng bộ này; số ở đây chỉ tính nhiệm vụ giao năm 2026. Đề nghị xác nhận có cần xem cả nhiệm vụ của các năm trước còn tồn không.'])

# ---------------- CHỈ TIÊU 4 — chỉ số 52
page_break()
doc.add_heading('Chỉ tiêu 4. Hồ sơ thủ tục hành chính, dịch vụ công (chỉ số 52)', level=2)
doc.add_heading('4.1. Thông tin chung', level=3)
info_table([
    ['Eform nguồn', FIELDS[3]['eform'].strip()],
    ['Một bản ghi là', 'Số liệu một quý của một phường/xã'],
    ['Cách lấy theo kỳ', '**Cộng dồn** các quý I, II, III năm 2026.'],
    ['Kỳ hiển thị', f'9 tháng năm 2026 · {NDV * 3} bản ghi ({NDV} phường/xã × 3 quý)'],
    ['Nguồn chính và cách ghi tên', 'Áp dụng từ 12/10/2026 (quy tắc hiển thị H19): số chính thức về hồ sơ thủ tục hành chính là số của Trung tâm Phục vụ hành chính công (hệ thống dịch vụ công). Số của chỉ tiêu này là số theo báo cáo eform của phường, xã, nên trên IOC các chỉ tiêu hồ sơ tiếp nhận và tỷ lệ giải quyết đúng hạn ghi kèm "(theo báo cáo eform)" để không nhầm với số của Trung tâm.'],
])
doc.add_heading('4.2. Trường dữ liệu sử dụng', level=3)
fields_table(ky('Lấy quý I – III; «Quý» còn là trục của biểu đồ 4.1') + [
    ['Đơn vị báo cáo', 'Danh mục', 'Nhóm', 'Tách Phường / Xã cho biểu đồ 4.3; các cột của bảng phường/xã'],
    ['Tổng số hồ sơ tiếp nhận', 'Số', 'Cộng', 'Thẻ số; mẫu số của tỷ lệ trực tuyến, số hóa; trọng số của điểm hài lòng'],
    ['Hồ sơ giải quyết đúng hạn', 'Số', 'Cộng', 'Tử số tỷ lệ đúng hạn; biểu đồ 4.1, 4.2'],
    ['Hồ sơ giải quyết quá hạn', 'Số', 'Cộng', 'Thẻ Quá hạn; biểu đồ 4.1, 4.2'],
    ['Hồ sơ trực tuyến', 'Số', 'Cộng', 'Tỷ lệ trực tuyến; mẫu số tỷ lệ thanh toán trực tuyến'],
    ['Hồ sơ thanh toán trực tuyến', 'Số', 'Cộng', 'Tỷ lệ thanh toán trực tuyến'],
    ['Hồ sơ đã số hóa', 'Số', 'Cộng', 'Tỷ lệ số hóa'],
    ['Điểm hài lòng (%)', 'Số', 'Cộng', 'Thẻ Mức độ hài lòng: bình quân có trọng số theo số hồ sơ tiếp nhận'],
], [])
doc.add_heading('4.3. Thẻ số tổng hợp', level=3)
kpi_table([
    ['Hồ sơ tiếp nhận', 'TỔNG(«Tổng số hồ sơ tiếp nhận»), 126 phường/xã × 3 quý', f'{n(h["tn"])} hồ sơ'],
    ['Giải quyết đúng hạn', '«Hồ sơ giải quyết đúng hạn» ÷ (đúng hạn + quá hạn) × 100', f'{n(h["dh"])} ÷ {n(h["dh"] + h["qh"])} = {pc(h["dh"], h["dh"] + h["qh"], 2)}'],
    ['Quá hạn', 'TỔNG(«Hồ sơ giải quyết quá hạn»)', f'{n(h["qh"])} hồ sơ'],
    ['Hồ sơ trực tuyến', 'TỔNG(«Hồ sơ trực tuyến») ÷ Hồ sơ tiếp nhận × 100', f'{n(h["tt"])} ÷ {n(h["tn"])} = {pc(h["tt"], h["tn"])}'],
    ['Thanh toán trực tuyến', 'TỔNG(«Hồ sơ thanh toán trực tuyến») ÷ TỔNG(«Hồ sơ trực tuyến») × 100', f'{n(h["tttt"])} ÷ {n(h["tt"])} = {pc(h["tttt"], h["tt"])}'],
    ['Đã số hóa', 'TỔNG(«Hồ sơ đã số hóa») ÷ Hồ sơ tiếp nhận × 100', f'{n(h["sh"])} ÷ {n(h["tn"])} = {pc(h["sh"], h["tn"])}'],
    ['Mức độ hài lòng', 'TỔNG(«Điểm hài lòng (%)» × «Tổng số hồ sơ tiếp nhận») ÷ TỔNG(«Tổng số hồ sơ tiếp nhận»); ghi %, 2 số lẻ như các tỷ lệ hồ sơ thủ tục hành chính', n(h['hl'], 2) + '%'],
])
doc.add_heading('4.4. Biểu đồ', level=3)
qy = s52['quy']
chart('4.1', 'Kết quả giải quyết theo quý', 'Biểu đồ cột xếp chồng',
      'Trục ngang: «Quý» I → III; mỗi cột chia 2 đoạn: Đúng hạn / Quá hạn', 'Số hồ sơ đã giải quyết; dòng phụ ghi số hồ sơ tiếp nhận của quý',
      ['Nhóm theo «Quý»:', 'Đúng hạn = TỔNG(«Hồ sơ giải quyết đúng hạn»); Quá hạn = TỔNG(«Hồ sơ giải quyết quá hạn»).',
       'Không vẽ đoạn "đang giải quyết" = tiếp nhận − đúng hạn − quá hạn: hồ sơ giải quyết trong quý có thể là hồ sơ kỳ trước chuyển sang mà eform không có trường này (bản trả lời BA, mục 06).'],
      ' · '.join(f'{QN[q]}: {n(o["dh"])} + {n(o["qh"])} = {n(o["dh"] + o["qh"])} (tiếp nhận {n(o["tn"])})' for q, o in enumerate(qy)))
chart('4.2', 'Kết quả giải quyết — 9 tháng', 'Biểu đồ tròn dạng vành khuyên; giữa vòng ghi số hồ sơ đã giải quyết',
      '2 phần: Giải quyết đúng hạn / Giải quyết quá hạn', 'Số hồ sơ; di chuột hiện thêm tỷ trọng %',
      ['Cộng 3 quý của biểu đồ 4.1; giữa vòng = đúng hạn + quá hạn.',
       'Dòng phụ "Chưa có kết quả trong kỳ" = TỔNG(«Tổng số hồ sơ tiếp nhận») − đúng hạn − quá hạn, chỉ hiện khi không âm; âm thì ghi "Có giải quyết hồ sơ kỳ trước" (bản trả lời BA, mục 06).'],
      f'Đúng hạn {n(h["dh"])} · Quá hạn {n(h["qh"])} → giữa vòng {n(h["dh"] + h["qh"])}; dòng phụ: chưa có kết quả trong kỳ {n(h["dangGq"])}')
cp = s52['cap']
def tl4(o): return [pc(o['dh'], o['dh'] + o['qh']), pc(o['tt'], o['tn']), pc(o['tttt'], o['tt']), pc(o['sh'], o['tn'])]
chart('4.3', 'Các tỷ lệ: phường và xã', 'Biểu đồ cột, mỗi tỷ lệ 2 cột đứng cạnh nhau (Phường / Xã)',
      'Trục ngang: 4 tỷ lệ Đúng hạn / Trực tuyến / Thanh toán trực tuyến / Đã số hóa', 'Tỷ lệ (%)',
      'Tách đơn vị theo chữ đầu tên trong «Đơn vị báo cáo» (Phường / Xã); trong mỗi nhóm cộng từng trường số rồi tính 4 tỷ lệ như ở thẻ số.',
      'Phường: ' + ' · '.join(tl4(cp[0])) + ' — Xã: ' + ' · '.join(tl4(cp[1])),
      f'Phường {n(cp[0]["tn"])} + xã {n(cp[1]["tn"])} = {n(h["tn"])} hồ sơ tiếp nhận.')
notes('4', ['Trung tâm Phục vụ hành chính công Thành phố đã có số hồ sơ lấy qua API. Số ở eform này để đối chiếu; khi hai nguồn lệch nhau cần thống nhất lấy nguồn nào làm số chính thức.',
            'Eform không có trường số hồ sơ phải nộp phí, nên tỷ lệ thanh toán trực tuyến đang tính trên số hồ sơ trực tuyến. Nếu cần tính trên hồ sơ có phí, đề nghị bổ sung trường này.',
            'Eform không có «Hồ sơ kỳ trước chuyển sang» và «Hồ sơ đang giải quyết cuối kỳ». Đề nghị bổ sung theo cấu trúc báo cáo tình hình giải quyết thủ tục hành chính (kỳ trước chuyển sang + tiếp nhận mới; đã giải quyết; đang giải quyết); khi có hai trường này, vòng kết quả trở lại 3 phần.',
            '«Điểm hài lòng (%)»: IOC theo tên trường, ghi là tỷ lệ % và gọi "Mức độ hài lòng". Đề nghị xác nhận là tỷ lệ % hay điểm trên thang 100; nếu là điểm, IOC sẽ ghi "điểm" thay cho "%".',
            'Dữ liệu mẫu trong file mô tả: «Hồ sơ thanh toán trực tuyến» (87.845.454) và «Hồ sơ đã số hóa» (42.457.454.584) lớn hơn «Tổng số hồ sơ tiếp nhận» (123.456), cần sửa lại.',
            'Eform không có «Cấp báo cáo»; các sở, ngành không báo qua eform này. Đề nghị xác nhận.'])

# ---------------- CHỈ TIÊU 5 — chỉ số 11, 12
page_break()
doc.add_heading('Chỉ tiêu 5. Hạ tầng công nghệ thông tin (chỉ số 11, 12)', level=2)
doc.add_heading('5.1. Thông tin chung', level=3)
du_tong = ssum(s11['dapUng'])
info_table([
    ['Eform nguồn', FIELDS[4]['eform'].strip()],
    ['Một bản ghi là', 'Một loại hạ tầng («Loại hạ tầng CNTT») của một đơn vị, ở một «Tình trạng hoạt động», kèm «Số lượng». Mức độ đáp ứng đếm một lượt cho mỗi cặp đơn vị × loại hạ tầng'],
    ['Cách lấy theo kỳ', '**Bản của kỳ đang xem.** Số lượng là số đang có, nên chỉ lấy bản báo cáo quý mới nhất, không cộng các quý.'],
    ['Kỳ hiển thị', f'Quý III/2026 · {n(du_tong)} lượt đánh giá (đơn vị × loại hạ tầng)'],
])
doc.add_heading('5.2. Trường dữ liệu sử dụng', level=3)
fields_table(ky('Chọn bản báo cáo của quý đang xem') + [
    ['Đơn vị báo cáo', 'Danh mục', 'Nhóm', 'Cột "Phòng họp có giao ban" của bảng phường/xã'],
    ['Cấp báo cáo', 'Danh mục', 'Nhóm', 'Tách Phường/xã / Sở/ngành: biểu đồ 5.1, 5.3, 5.4'],
    ['Loại hạ tầng CNTT', 'Danh mục', 'Nhóm', 'Trục của biểu đồ 5.1; tách riêng Máy tính làm việc ra thẻ số'],
    ['Tình trạng hoạt động', 'Danh mục', 'Điều kiện, Nhóm', 'Thẻ Đang hoạt động, Hỏng; biểu đồ 5.2'],
    ['Số lượng', 'Số', 'Cộng', 'Mọi con số về thiết bị, hệ thống'],
    ['Mức độ đáp ứng', 'Danh mục', 'Nhóm', 'Thẻ Đáp ứng đầy đủ; biểu đồ 5.3'],
    ['Tổng số phòng họp (chỉ áp dụng HT07)', 'Số', 'Cộng', 'Mẫu số tỷ lệ phòng họp có giao ban trực tuyến'],
    ['Số phòng được trang bị hệ thống giao ban trực tuyến (HT07)', 'Số', 'Cộng', 'Tử số tỷ lệ trên (không dùng thay số hệ thống)'],
], ['Đối tượng phục vụ', 'Ghi chú'])
para('"HT07" trong tên hai trường phòng họp là mã của giá trị "Hệ thống giao ban trực tuyến" trong danh mục «Loại hạ tầng CNTT». Hai trường phòng họp là số của cả đơn vị: mỗi đơn vị, mỗi kỳ lấy một lần ở dòng HT07 (nhiều dòng thì lấy bản mới nhất), không cộng giữa các dòng; «Số lượng» của dòng HT07 là số hệ thống (bản trả lời BA, mục 04).', size=11.5, color=GREY, italic=True)
doc.add_heading('5.3. Số lượng theo loại hạ tầng', level=3)
para('Bảng dưới là TỔNG(«Số lượng») theo «Loại hạ tầng CNTT» và «Cấp báo cáo», cộng mọi tình trạng.').paragraph_format.keep_with_next = True
rows = [['Loại hạ tầng', 'Phường/xã', 'Sở/ngành', 'Cộng']]
for i, r in enumerate(s11['loai']): rows.append([r[0], n(ssum(r[1])), n(ssum(r[2])), n(s11['loaiTong'][i])])
rows.append(['**Cộng**', f'**{n(ssum(ssum(r[1]) for r in s11["loai"]))}**', f'**{n(ssum(ssum(r[2]) for r in s11["loai"]))}**', f'**{n(s11["tong"])}**'])
table(rows, [7.0, 3.0, 3.0, 3.0], size=11, keep=True)
doc.add_heading('5.4. Thẻ số tổng hợp', level=3)
pc_xa = ssum(s11['loai'][0][1]); pc_so = ssum(s11['loai'][0][2])
kpi_table([
    ['Tổng thiết bị, hệ thống', 'TỔNG(«Số lượng») mọi loại, mọi tình trạng, cả 2 cấp', n(s11['tong'])],
    ['Đang hoạt động', 'TỔNG(«Số lượng») có «Tình trạng hoạt động» = Đang hoạt động ÷ Tổng × 100', f'{n(s11["tt"][0])} ÷ {n(s11["tong"])} = {pc(s11["tt"][0], s11["tong"])}'],
    ['Hỏng/không hoạt động', 'TỔNG(«Số lượng») có «Tình trạng hoạt động» = Hỏng/Không hoạt động', n(s11['tt'][3])],
    ['Máy tính làm việc', 'TỔNG(«Số lượng») có «Loại hạ tầng CNTT» = Máy tính làm việc', f'{n(pc_xa)} + {n(pc_so)} = {n(s11["loaiTong"][0])}'],
    ['Phòng họp có giao ban trực tuyến', 'TỔNG(«Số phòng được trang bị…») ÷ TỔNG(«Tổng số phòng họp…») × 100', f'{n(s11["phongTb"])} ÷ {n(s11["phong"])} = {pc(s11["phongTb"], s11["phong"])}'],
    ['Đáp ứng đầy đủ', 'ĐẾM lượt «Mức độ đáp ứng» = Đáp ứng đầy đủ ÷ tổng số lượt đánh giá × 100', f'{n(s11["dapUng"][0])} ÷ {n(du_tong)} = {pc(s11["dapUng"][0], du_tong)}'],
])
doc.add_heading('5.5. Biểu đồ', level=3)
L = s11['loai']
chart('5.1', 'Số lượng theo loại hạ tầng (trừ máy tính)', 'Biểu đồ cột ngang xếp chồng',
      'Trục dọc: 9 giá trị của «Loại hạ tầng CNTT» (trừ Máy tính làm việc); mỗi cột chia 2 đoạn theo «Cấp báo cáo»',
      'Số lượng; số ở đầu cột = cộng 2 cấp',
      'Nhóm theo «Loại hạ tầng CNTT» và «Cấp báo cáo», mỗi đoạn = TỔNG(«Số lượng»).',
      f'{L[2][0]}: {n(ssum(L[2][1]))} + {n(ssum(L[2][2]))} = {n(s11["loaiTong"][2])} · {L[1][0]}: {n(ssum(L[1][1]))} + {n(ssum(L[1][2]))} = {n(s11["loaiTong"][1])}',
      f'Máy tính làm việc ({n(s11["loaiTong"][0])}) lớn hơn hẳn các loại khác nên hiện ở thẻ số riêng; nếu vẽ chung, các loại còn lại gần như không nhìn thấy.')
TT = [t[0] for t in d11['tt']]
chart('5.2', 'Theo tình trạng hoạt động', 'Biểu đồ tròn dạng vành khuyên; giữa vòng ghi tổng thiết bị, hệ thống',
      '4 giá trị của «Tình trạng hoạt động»: ' + ' / '.join(TT), 'Số lượng; di chuột hiện thêm tỷ trọng %',
      'Nhóm theo «Tình trạng hoạt động», mỗi lát = TỔNG(«Số lượng»), mọi loại hạ tầng, cả 2 cấp.',
      ' · '.join(f'{TT[i]} {n(s11["tt"][i])}' for i in range(4)) + f' → giữa vòng {n(s11["tong"])}')
DU = [t[0] for t in d11['dapUng']]
chart('5.3', 'Mức độ đáp ứng theo cấp', 'Biểu đồ cột ngang xếp chồng',
      'Trục dọc: «Cấp báo cáo» (Phường/xã, Sở/ngành); mỗi cột chia 3 đoạn: ' + ' / '.join(DU), 'Số lượt đánh giá (đơn vị × loại hạ tầng)',
      ['Nhóm theo «Cấp báo cáo», trong mỗi cấp ĐẾM số lượt theo «Mức độ đáp ứng».', 'Mỗi cặp đơn vị × loại hạ tầng một lượt, không đếm theo dòng tình trạng: nhiều dòng thì lấy bản mới nhất; cùng thời điểm mà mức khác nhau thì lấy mức thấp nhất (bản trả lời BA, mục 03).'],
      f'Phường/xã: ' + ' + '.join(n(v) for v in d11['dapUngXa']) + f' = {n(ssum(d11["dapUngXa"]))} · Sở/ngành: ' + ' + '.join(n(v) for v in d11['dapUngSo']) + f' = {n(ssum(d11["dapUngSo"]))}')
chart('5.4', 'Phòng họp có hệ thống giao ban trực tuyến', 'Biểu đồ cột ngang, mỗi cấp 2 cột đứng cạnh nhau',
      'Trục dọc: «Cấp báo cáo»; 2 cột: Tổng số phòng họp / Có giao ban trực tuyến', 'Số phòng',
      'Nhóm theo «Cấp báo cáo»: cột 1 = TỔNG(«Tổng số phòng họp (chỉ áp dụng HT07)»); cột 2 = TỔNG(«Số phòng được trang bị hệ thống giao ban trực tuyến (HT07)»).',
      f'Phường/xã: {n(s11["tbXa"])} / {n(s11["phXa"])} ({pc(s11["tbXa"], s11["phXa"])}) · Sở/ngành: {n(d11["soPhongTb"])} / {n(d11["soPhongHop"])} ({pc(d11["soPhongTb"], d11["soPhongHop"])})')
notes('5', ['Một đơn vị có thể khai nhiều bản ghi cùng một loại hạ tầng, mỗi bản ghi một «Tình trạng hoạt động»; IOC cộng «Số lượng» theo từng tình trạng. Đề nghị xác nhận cách nhập này.',
            '«Số lượng» của Hệ thống giao ban trực tuyến là số hệ thống, không dùng thay số phòng (bản trả lời BA, mục 04). Số giả lập: phường, xã có phòng được trang bị 1 hệ thống (phường trung tâm 2), sở, ngành 92 hệ thống. Hai dòng HT07 của cùng đơn vị ghi số phòng khác nhau thì IOC gắn cờ; số phòng có giao ban lớn hơn tổng số phòng thì gắn cờ, không vẽ.',
            'Đề nghị Văn phòng xác nhận cụ thể chỉ số 11 và chỉ số 12 là con số nào trong các thẻ số trên (IOC đang giả định chỉ số 12 là tỷ lệ phòng họp có giao ban trực tuyến).',
            '«Mức độ đáp ứng» cần ghi rõ đáp ứng so với yêu cầu nào để các đơn vị đánh giá thống nhất.',
            'Dữ liệu mẫu: «Số phòng được trang bị…» (999.999.999) lớn hơn «Tổng số phòng họp…» (123.456), cần sửa lại.'])

# ======================================================= PHẦN III
page_break()
doc.add_heading('PHẦN III. BẢNG 126 PHƯỜNG/XÃ', level=1)
para(f'Ô bên phải của form là bảng {NDV} phường/xã. Dòng đầu **Cộng {NDV} phường/xã** chỉ cộng các phường/xã, **chưa gồm các sở, ban, ngành**. Ở cột tỷ lệ, dòng đầu tính từ tổng tử số và mẫu số, không lấy trung bình các dòng (Phần I, mục 5).')
nv_db = d74['coDiaBan']
table([
    ['Cột trong bảng', 'Chỉ tiêu', 'Nhóm theo trường', 'Cách tính', 'Dòng đầu'],
    ['VB xử lý trên mạng quý III (%)', '1', '«Đơn vị báo cáo»', 'Số xử lý trên mạng ÷ Tổng số văn bản, hồ sơ × 100, quý III', pc(s27['xaMang'][Q3], s27['xaTong'][Q3])],
    ['Trợ lý ảo', '2', '«Đơn vị báo cáo»', 'Có / Chưa theo «Đã triển khai trợ lý ảo»', f'{s29["tla"]} Có'],
    ['Hồ sơ tiếp nhận 9 tháng', '4', '«Đơn vị báo cáo»', 'TỔNG(«Tổng số hồ sơ tiếp nhận») quý I – III', n(h['tn'])],
    ['Đúng hạn (%)', '4', '«Đơn vị báo cáo»', 'Đúng hạn ÷ (đúng hạn + quá hạn) × 100, quý I – III', pc(h['dh'], h['dh'] + h['qh'])],
    ['Phòng họp có giao ban', '5', '«Đơn vị báo cáo»', 'Số phòng được trang bị / Tổng số phòng họp', f'{n(s11["tbXa"])}/{n(s11["phXa"])}'],
    ['Nhiệm vụ CP, TTg có địa bàn', '3', '«Địa bàn»', 'ĐẾM(«Mã nhiệm vụ»); nhiệm vụ chung của Thành phố (không chọn địa bàn) không vào bảng', f'{n(nv_db)} / {n(s74["tong"])}'],
], [3.7, 1.5, 3.0, 5.8, 2.0], size=10.5)
u = U[9]; uh = [sum(x[k] for x in u['hs']) for k in ('tn', 'dh', 'qh')]
box([f'**{u["ten"]}**: văn bản xử lý trên mạng quý III {n(u["vbMang"][Q3])} ÷ {n(u["vbTong"][Q3])} = {pc(u["vbMang"][Q3], u["vbTong"][Q3])} · trợ lý ảo {"Có" if u["tla"] else "Chưa"} · '
     f'hồ sơ tiếp nhận {" + ".join(n(x["tn"]) for x in u["hs"])} = {n(uh[0])} · đúng hạn {n(uh[1])} ÷ {n(uh[1] + uh[2])} = {pc(uh[1], uh[1] + uh[2])} · phòng họp {u["phongTb"]}/{u["phongHop"]}.'],
    FORM_FILL, 'Ví dụ một dòng: ')
para('Bấm vào một biểu đồ bất kỳ ở cột trái thì ô bên phải đổi sang bảng số của chính biểu đồ đó; bấm nút "Danh sách 126 đơn vị" để quay lại bảng phường/xã.', italic=True, color=GREY, size=12)

# ======================================================= PHẦN IV
page_break()
doc.add_heading('PHẦN IV. TỔNG HỢP ĐỀ NGHỊ VĂN PHÒNG XÁC NHẬN', level=1)
para('Gom các đề nghị Văn phòng xác nhận của 5 chỉ tiêu. Cột "Kết quả" để trống cho Văn phòng ghi ý kiến.')
table([['#', 'Chỉ tiêu', 'Nội dung', 'Kết quả']] + [[str(i + 1), c, x, ''] for i, (c, x) in enumerate(ALL_NOTES)],
      [0.8, 1.6, 10.6, 3.0], size=10.5)

# ======================================================= PHỤ LỤC
page_break()
ctm_section(lambda t, l: doc.add_heading(t, level=l), para, lambda rows, w: table(rows, w, size=10), 'ubnd-vpubnd',
            'PHẦN V. CÁCH HIỂN THỊ TRÊN GIAO DIỆN IOC', changes=['Chỉ tiêu hồ sơ thủ tục hành chính ghi rõ **"theo báo cáo eform"**; số chính thức về hồ sơ là của Trung tâm Phục vụ hành chính công.', 'Bảng 126 phường/xã: văn bản xử lý trên môi trường mạng và phòng họp có giao ban dùng **"Mức chung 126 phường, xã"** (không gồm sở, ban, ngành).', 'Tỷ lệ giải quyết đúng hạn và mức độ hài lòng ghi **2 số lẻ**. Trường "Điểm hài lòng (%)" hiện là **"Mức độ hài lòng"** (%), không ghi "điểm … %"; chờ Văn phòng xác nhận là tỷ lệ hay điểm thang 100.', 'Mức độ đáp ứng hạ tầng đếm **một lượt cho mỗi cặp đơn vị × loại hạ tầng**, ghi "lượt đánh giá" (BA mục 03).', '«Số lượng» của Hệ thống giao ban trực tuyến là **số hệ thống**, không lấy bằng số phòng được trang bị (BA mục 04).', 'Kết quả giải quyết hồ sơ: vòng **2 phần** Đúng hạn / Quá hạn, giữa vòng = số đã giải quyết; phần chênh ghi dòng phụ "Chưa có kết quả trong kỳ" (BA mục 06).', 'Thẻ ở màn ngoài dùng một bộ nhãn rõ nghĩa, giống nhau ở tab UBND và tab "Tất cả"; tab UBND đọc đúng số của màn chi tiết, không tự tính lại.', 'Mọi mũi tên tăng/giảm ghi rõ mốc so sánh; chưa có số cùng kỳ thì ghi "Chưa có số cùng kỳ" thay cho mũi tên.', 'Tỷ lệ theo tháng/quý vẽ đường thay cho cột từ 0; bảng cơ cấu tối đa 6 phần chuyển sang thẻ cơ cấu; bỏ thẻ lặp đúng bộ số của thẻ khác.'])
page_break()
page_break()
doc.add_heading('PHỤ LỤC. DANH SÁCH TRƯỜNG CỦA 5 EFORM', level=1)
para('Trích nguyên từ file mô tả đầu vào eform, sheet "Văn phòng UBND Thành phố Hà Nội". Cột "Dùng trên IOC" cho biết trường có tham gia tính toán biểu đồ / thẻ số hay không.')
KY = {'Kỳ báo cáo', 'Năm', 'Quý', 'Tháng'}
USED = {
    0: KY | {'Đơn vị báo cáo', 'Cấp báo cáo', 'Tổng số văn bản, hồ sơ công việc', 'Số văn bản, hồ sơ xử lý trên môi trường mạng'},
    1: KY | {'Đơn vị báo cáo', 'Đối tượng phục vụ', 'Đã triển khai trợ lý ảo', 'Tên ứng dụng trợ lý ảo'},
    2: KY | {'Mã nhiệm vụ', 'Cơ quan giao', 'Ngày giao', 'Hạn hoàn thành', 'Đơn vị chủ trì', 'Trạng thái thực hiện', 'Tỷ lệ hoàn thành (%)', 'Ngày hoàn thành', 'Cảnh báo tiến độ', 'Địa bàn'},
    3: KY | {'Đơn vị báo cáo', 'Tổng số hồ sơ tiếp nhận', 'Hồ sơ giải quyết đúng hạn', 'Hồ sơ giải quyết quá hạn', 'Hồ sơ trực tuyến', 'Hồ sơ thanh toán trực tuyến', 'Hồ sơ đã số hóa', 'Điểm hài lòng (%)'},
    4: KY | {'Đơn vị báo cáo', 'Cấp báo cáo', 'Loại hạ tầng CNTT', 'Tình trạng hoạt động', 'Số lượng', 'Mức độ đáp ứng', 'Tổng số phòng họp (chỉ áp dụng HT07)', 'Số phòng được trang bị hệ thống giao ban trực tuyến (HT07)'},
}
used_total = 0
for i, e in enumerate(FIELDS):
    doc.add_heading(f'{i + 1}. {e["eform"].strip()} ({len(e["fields"])} trường)', level=2)
    rows = [['#', 'Tên trường', 'Mã trường', 'Kiểu nhập', 'Dùng trên IOC']]
    fills = {}
    for j, f in enumerate(e['fields']):
        u_ = f[0] in USED[i]; used_total += u_
        rows.append([str(j + 1), f[0], f[1], f[2] or '(chưa ghi)', 'Có' if u_ else '—'])
        if u_: fills[(j + 1, 4)] = 'E6F4EA'
    t = table(rows, [0.8, 4.9, 6.5, 2.0, 1.8], size=10, fills=fills)
    for j in range(1, len(rows)):  # mã trường dùng phông đơn cách
        for r in t.cell(j, 2).paragraphs[0].runs: r.font.name = 'Consolas'; r.font.size = Pt(8.5)
for i, us in USED.items():  # tên trường trong USED phải có thật trong eform
    missing = us - {f[0] for f in FIELDS[i]['fields']}
    assert not missing, (i, missing)
print('Trường dùng trên IOC:', used_total, '/', N_FIELDS, '· đề nghị:', len(ALL_NOTES))

os.makedirs(OUT_DIR, exist_ok=True)
doc.save(OUT)
MD.save(OUT_MD)
print('ok', OUT)
print('ok', OUT_MD)
