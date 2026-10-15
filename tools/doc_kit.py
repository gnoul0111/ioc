# -*- coding: utf-8 -*-
"""Bộ dựng tài liệu mô tả chỉ tiêu dùng chung: mỗi lệnh ghi ĐỒNG THỜI vào Word (.docx) và Markdown (.md).

Định dạng Word giống tools/vpub_doc_docx.py và tools/vhtt_doc_docx.py (A4 lề 3-2-2-2, Times New Roman 13, bảng tiêu đề
nền xanh đậm, khung ghi chú nền vàng nhạt, mục lục tự cập nhật bằng tools/vhtt_doc_finalize.ps1). Phần Markdown do
tools/doc_md.py lo. Tài liệu mới dùng module này để khỏi chép lại ~200 dòng tiện ích.

    from doc_kit import Kit
    k = Kit('Đơn vị · Mô tả chỉ tiêu trên IOC')
    k.heading('PHẦN I. …', 1); k.para('…'); k.table([[…], …], [3, 13]); k.box([…], k.NOTE_FILL, 'Lưu ý: ')
    k.save('docs/…/X.docx')            # ghi cả X.docx và X.md

Chữ trong para/table/box: **đậm**, «tên trường» (in đậm nghiêng xanh trong Word, giữ nguyên trong Markdown)."""
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from doc_md import MdMirror

CENTER = WD_ALIGN_PARAGRAPH.CENTER
FONT = 'Times New Roman'
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2A, 0x5C, 0x9A)
GREY = RGBColor(0x55, 0x5F, 0x6D)
BORDER = 'A9B8CC'
PAGE_W = 16.0  # cm vùng chữ (A4, lề trái 3, phải 2)


def _set_font(style, size, bold=None, color=None, italic=None):
    style.font.name = FONT; style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'): rf.set(qn(a), FONT)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    if bold is not None: style.font.bold = bold
    if italic is not None: style.font.italic = italic
    if color is not None: style.font.color.rgb = color


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill); tcPr.append(s)


def _cell_margins(table, top=70, bottom=70, left=100, right=100):
    tblPr = table._tbl.tblPr; m = OxmlElement('w:tblCellMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement('w:' + k); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tblPr.append(m)


def _borders(table, color=BORDER):
    tblPr = table._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for k in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + k); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), color); b.append(e)
    tblPr.append(b)


def _fixed_layout(table, widths):
    tblPr = table._tbl.tblPr; lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(widths) * 567))); tw.set(qn('w:type'), 'dxa'); tblPr.append(tw)
    for i, gc in enumerate(table._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'), str(int(widths[i] * 567)))
    for row in table.rows:
        for i, c in enumerate(row.cells): c.width = Cm(widths[i])


def _row_flag(row, tag):
    trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement(tag))


def _field(p, instr):
    r = p.add_run(); b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._r.append(b)
    r = p.add_run(); t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = instr; r._r.append(t)
    r = p.add_run(); s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate'); r._r.append(s)
    r = p.add_run('…'); e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r2 = p.add_run(); r2._r.append(e)
    return r


def _bottom_rule(p, color='1F3A5F', sz='8'):
    pPr = p._p.get_or_add_pPr(); bd = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), sz); b.set(qn('w:space'), '4'); b.set(qn('w:color'), color); bd.append(b); pPr.append(bd)


def _add_runs(p, text, size=None, bold=False, color=None, italic=False):
    """**đậm**, «tên trường» in nghiêng đậm màu xanh."""
    for part in re.split(r'(\*\*[^*]+\*\*|«[^»]+»)', text):
        if not part: continue
        if part.startswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('«'):
            r = p.add_run(part); r.bold = True; r.italic = True; r.font.color.rgb = BLUE
        else:
            r = p.add_run(part); r.bold = bold; r.italic = italic
            if color is not None: r.font.color.rgb = color
        if size: r.font.size = Pt(size)
    return p


class Kit:
    HEAD_FILL, KEY_FILL, NOTE_FILL, FORM_FILL, OK_FILL = '1F3A5F', 'EEF3F9', 'FFF7E6', 'F3F6FA', 'E6F4EA'
    GREY = GREY

    def __init__(self, header_text):
        self.doc = doc = Document(); self.md = MdMirror(); self._pb = False
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(2), Cm(2), Cm(2)
        sec.header_distance, sec.footer_distance = Cm(1.1), Cm(1.1)
        st = doc.styles
        _set_font(st['Normal'], 13)
        pf = st['Normal'].paragraph_format; pf.space_after = Pt(6); pf.line_spacing = 1.2
        for name, size, before, after, color in (('Heading 1', 15, 18, 8, NAVY), ('Heading 2', 13.5, 14, 6, NAVY), ('Heading 3', 13, 10, 4, BLUE)):
            s = st[name]; _set_font(s, size, bold=True, color=color, italic=False)
            s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.keep_with_next = True; s.paragraph_format.line_spacing = 1.15
        for name in ('TOC 1', 'TOC 2', 'TOC 3'):
            try: s = st[name]
            except KeyError: s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            _set_font(s, 12.5 if name == 'TOC 1' else 12, bold=(name == 'TOC 1'))
            s.paragraph_format.space_after = Pt(3); s.paragraph_format.left_indent = Cm({'TOC 1': 0, 'TOC 2': 0.6, 'TOC 3': 1.2}[name])
        _set_font(st['List Bullet'], 13); st['List Bullet'].paragraph_format.space_after = Pt(3)
        # đầu trang, chân trang (trang bìa không có)
        hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(header_text); r.italic = True; r.font.size = Pt(10); r.font.color.rgb = GREY
        _bottom_rule(hp, 'A9B8CC', '4')
        fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run('Trang '); r.font.size = Pt(10); r.font.color.rgb = GREY
        r = _field(fp, 'PAGE'); r.font.size = Pt(10)
        r = fp.add_run(' / '); r.font.size = Pt(10); r.font.color.rgb = GREY
        r = _field(fp, 'NUMPAGES'); r.font.size = Pt(10)
        sec.different_first_page_header_footer = True

    # ---- sang trang: gắn page_break_before vào đoạn kế tiếp (tránh trang trắng) ----
    def page_break(self): self._pb = True

    def _take_pb(self, p):
        if self._pb: p.paragraph_format.page_break_before = True; self._pb = False
        return p

    def _gap(self):
        g = self.doc.add_paragraph(); g.paragraph_format.space_after = Pt(4); g.paragraph_format.line_spacing = 0.8

    # ---- khối nội dung ----
    def blank(self, n=1):
        for _ in range(n): self.doc.add_paragraph()

    def heading(self, text, level=1):
        self.md.heading(text, level)
        return self._take_pb(self.doc.add_heading(text, level=level))

    def para(self, text, size=None, align=None, after=None, bold=False, color=None, italic=False, md=True):
        p = self._take_pb(self.doc.add_paragraph()); _add_runs(p, text, size, bold, color, italic)
        if md: self.md.para(text, size, italic, bold)
        if align is not None: p.alignment = align
        if after is not None: p.paragraph_format.space_after = Pt(after)
        return p

    def note(self, text):
        """Đoạn chú thích nhỏ, nghiêng, xám."""
        return self.para(text, 11.5, italic=True, color=GREY)

    def bullets(self, items, size=None):
        self.md.bullets(items)
        for t in items:
            p = self._take_pb(self.doc.add_paragraph(style='List Bullet')); _add_runs(p, t, size)

    def rule(self, indent=5.5):
        p = self.doc.add_paragraph(); p.alignment = CENTER; p.paragraph_format.space_after = Pt(0)
        _bottom_rule(p, '1F3A5F', '6'); p.paragraph_format.left_indent = Cm(indent); p.paragraph_format.right_indent = Cm(indent)

    def toc(self, title='MỤC LỤC'):
        self.para(title, 15, CENTER, 10, bold=True, color=NAVY)
        _field(self.doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u'); self.md.toc()

    def table(self, rows, widths, header=True, size=11, key_col=False, fills=None, keep=False, mono_col=None):
        self.md.table(rows, header)
        t = self.doc.add_table(rows=len(rows), cols=len(widths)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t); _cell_margins(t); _fixed_layout(t, widths)
        for ri, row in enumerate(rows):
            _row_flag(t.rows[ri], 'w:cantSplit')
            for ci, val in enumerate(row):
                c = t.cell(ri, ci); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
                for li, ln in enumerate(val if isinstance(val, list) else [val]):
                    if li: p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
                    if header and ri == 0:
                        r = p.add_run(ln); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    else:
                        _add_runs(p, ln, size, bold=(key_col and ci == 0))
                        if mono_col == ci and ri:
                            for r in p.runs: r.font.name = 'Consolas'; r.font.size = Pt(8.5)
                if header and ri == 0: _shade(c, self.HEAD_FILL)
                elif key_col and ci == 0: _shade(c, self.KEY_FILL)
                elif fills and fills.get((ri, ci)): _shade(c, fills[(ri, ci)])
        if header: _row_flag(t.rows[0], 'w:tblHeader')
        if keep:  # giữ cả bảng trên một trang
            for row in t.rows[:-1]:
                for c in row.cells:
                    for p in c.paragraphs: p.paragraph_format.keep_with_next = True
        self._gap()
        return t

    def box(self, lines, fill=None, label=None, size=12):
        """Khung 1 ô có nền — dùng cho ví dụ / lưu ý."""
        self.md.box(lines, label)
        t = self.doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t, 'C9D4E3'); _cell_margins(t, 110, 110, 160, 160); _fixed_layout(t, [PAGE_W])
        _row_flag(t.rows[0], 'w:cantSplit')
        c = t.cell(0, 0); _shade(c, fill or self.FORM_FILL); p = c.paragraphs[0]; first = True
        if label:
            r = p.add_run(label); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = NAVY; first = False
        for ln in lines:
            if not first: p = c.add_paragraph()
            p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.15
            _add_runs(p, ln, size); first = False
        self._gap()

    # ---- ghi file ----
    def save(self, docx_path):
        os.makedirs(os.path.dirname(docx_path), exist_ok=True)
        self.doc.save(docx_path)
        md_path = os.path.splitext(docx_path)[0] + '.md'
        self.md.save(md_path)
        return docx_path, md_path
