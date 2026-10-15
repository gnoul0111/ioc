import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, dst = sys.argv[1], sys.argv[2]
lines = open(src, encoding="utf-8").read().split("\n")
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(2))
sec.top_margin = sec.bottom_margin = Cm(2)
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
HEAD = RGBColor(0x1F, 0x3A, 0x5F)

def runs(par, text, size=None):
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part: continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        else:
            r = par.add_run(part)
        if size and not part.startswith("`"): r.font.size = Pt(size)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), color); tcPr.append(s)

def heading(text, lvl):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.color.rgb = HEAD
    r.font.size = Pt({1: 18, 2: 14, 3: 12.5, 4: 12}[lvl]); p.paragraph_format.space_before = Pt(14 if lvl < 3 else 10); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True

i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("|"):
        block = []
        while i < len(lines) and lines[i].startswith("|"):
            block.append(lines[i]); i += 1
        rows = [[c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", r)[1:-1]] for r in block if not re.match(r"^\|[-| ]+\|$", r)]
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci); cell.text = ""
                runs(cell.paragraphs[0], val, 10)
                if ri == 0:
                    shade(cell, "DCE6F1")
                    for r in cell.paragraphs[0].runs: r.bold = True
        doc.add_paragraph()
        continue
    m = re.match(r"^(#{1,4}) (.*)", ln)
    if m: heading(m.group(2), len(m.group(1)))
    elif ln.startswith("- "):
        p = doc.add_paragraph(style="List Bullet"); runs(p, ln[2:])
    elif re.match(r"^\s+- ", ln):
        p = doc.add_paragraph(style="List Bullet 2"); runs(p, ln.lstrip()[2:])
    elif re.match(r"^\s*\d+\. ", ln):
        p = doc.add_paragraph(); runs(p, ln.strip())
        p.paragraph_format.left_indent = Cm(1.27 if ln.startswith(" ") else 0.63)
        p.paragraph_format.space_after = Pt(2)
    elif ln.strip():
        p = doc.add_paragraph(); runs(p, ln.strip()); p.paragraph_format.space_after = Pt(6)
        if ln.startswith("  "): p.paragraph_format.left_indent = Cm(0.63)
    i += 1
doc.save(dst)
print("ok", dst)
